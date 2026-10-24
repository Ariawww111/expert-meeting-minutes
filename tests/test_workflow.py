from __future__ import annotations

import copy
import sys
import tempfile
import unittest
from pathlib import Path


SKILL_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = SKILL_ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from common import (  # noqa: E402
    SCHEMA_VERSION,
    extract_numeric_tokens,
    normalize_numeric_token,
    read_json,
    sha256_json,
    summary_character_count,
    write_json,
)
from audio_pipeline import (  # noqa: E402
    _transcription_config,
    _parse_sentences,
    _write_run_manifest,
    _coverage_gaps,
    chunk_count,
    data_root_writable,
    parse_channels,
)
from docx_renderer import RenderError, render_docx  # noqa: E402
from fidelity import validate_minutes  # noqa: E402
from normalization import (  # noqa: E402
    ambiguous_term_matches,
    apply_proposal_markdown,
    load_rules,
    normalize_text,
)
from review_pipeline import _review_segment_issue, prepare_review  # noqa: E402


def _pad_summary(paragraphs, target=1100):
    result = list(paragraphs)
    filler = "上述讨论均围绕会议中已经出现的事实、条件和数据展开。"
    while summary_character_count(result) < target:
        result.append(filler)
        if len(result) > 5:
            result = [result[0], result[1], result[2], "".join(result[3:])]
    return result


def _fixture():
    transcript = {
        "schema_version": SCHEMA_VERSION,
        "audio": {
            "sha256": "a" * 64,
            "duration_seconds": 120.0,
            "selected_channel": "0",
        },
        "engine": {"name": "FunASR", "model": "paraformer-zh"},
        "coverage": {
            "status": "pass",
            "speech_seconds": 100.0,
            "transcribed_seconds": 100.0,
            "coverage_ratio": 1.0,
            "gaps": [],
        },
        "segments": [
            {
                "id": "S000001",
                "start": 0.0,
                "end": 12.0,
                "speaker": "C01-SPK1",
                "text": "专家表示行业需求增长20%，明年产能约10万吨。",
            },
            {
                "id": "S000002",
                "start": 12.0,
                "end": 18.0,
                "speaker": "C01-SPK2",
                "text": "与会人员询问价格走势。",
            },
            {
                "id": "S000003",
                "start": 18.0,
                "end": 28.0,
                "speaker": "C01-SPK1",
                "text": "专家认为价格可能保持稳定。",
            },
            {
                "id": "S000004",
                "start": 28.0,
                "end": 31.0,
                "speaker": "C01-SPK2",
                "text": "谢谢大家。",
            },
        ],
    }
    summary = _pad_summary(
        [
            "专家在会议中讨论了行业需求和产能变化。与会专家表示，行业需求增长20%，明年产能约10万吨。相关判断以会议中提出的条件和数据为基础，没有引入会议之外的信息。",
            "围绕价格走势，专家认为价格可能保持稳定。该表述保留了专家原有的不确定性，没有被改写为确定结论。专家同时说明了供需和产能条件对判断的影响，相关内容均来自会议现场。",
            "会议讨论集中于行业运行、供需变化和企业经营口径。所有数字、单位、时间条件和专家判断均按原始讨论保留，未增加投资建议、市场预期比较或纪要作者的独立评价。",
        ]
    )
    minutes = {
        "schema_version": SCHEMA_VERSION,
        "title": "专家会议纪要",
        "summary": summary,
        "summary_segment_ids": ["S000001"],
        "qa": [
            {
                "id": "Q1",
                "question": "价格走势如何判断？",
                "answer": ["价格可能保持稳定。"],
                "evidence_segment_ids": ["S000002", "S000003"],
            }
        ],
        "omissions": [{"segment_id": "S000004", "reason": "greeting"}],
    }
    review = {
        "schema_version": SCHEMA_VERSION,
        "transcript_sha256": sha256_json(transcript["segments"]),
        "status": "ready",
        "issues": [
            {
                "id": "RTEST001",
                "kind": "number",
                "kinds": ["number"],
                "segment_ids": ["S000001"],
                "time_range": [0.0, 12.0],
                "quote": transcript["segments"][0]["text"],
                "candidates": ["20%", "10万吨"],
                "reason": "数字或单位需要人工确认",
                "source_checks": [],
                "status": "confirmed",
                "replacement_text": transcript["segments"][0]["text"],
                "note": "测试确认",
            }
        ],
    }
    glossary = {"schema_version": SCHEMA_VERSION, "terms": [], "person_names": ["张三"]}
    return transcript, minutes, review, glossary


class WorkflowTests(unittest.TestCase):
    def test_audio_metadata_helpers(self):
        self.assertEqual(parse_channels("mono"), 1)
        self.assertEqual(parse_channels("stereo"), 2)
        self.assertEqual(parse_channels("5.1(side)"), 6)
        self.assertEqual(chunk_count(600.0, 600), 1)
        self.assertEqual(chunk_count(600.1, 600), 2)
        self.assertEqual(chunk_count(3561.19, 600), 6)
        with tempfile.TemporaryDirectory() as temp:
            self.assertTrue(data_root_writable(Path(temp) / "data"))

    def test_numeric_tokens_work_inside_chinese_sentences(self):
        tokens = extract_numeric_tokens(
            "上半年超过1.5亿元，销量一点五亿片，随访3个月，汇兑亏损3000多万元。"
        )
        self.assertIn("1.5亿元", tokens)
        self.assertIn("一点五亿片", tokens)
        self.assertIn("3个月", tokens)
        self.assertIn("3000多万元", tokens)

    def test_safe_normalization_handles_dates_and_standard_abbreviations(self):
        rules = load_rules()
        text, changes = normalize_text(
            "专利延长到三三年，四五六三个月，五六月份，七月二十号，"
            "PKP暴露良好，LPA靶点项目。",
            rules,
        )
        self.assertIn("2033年", text)
        self.assertIn("4月至6月", text)
        self.assertIn("5月至6月", text)
        self.assertIn("7月20日", text)
        self.assertIn("PK/PD", text)
        self.assertIn("Lp(a)", text)
        self.assertTrue(changes)

    def test_safe_normalization_removes_duplicate_month_fragments(self):
        rules = load_rules()
        text, _ = normalize_text(
            "公告是七月七月二号发的，十月底左右完成，十月十月底读出数据。",
            rules,
        )
        self.assertEqual(text, "公告是7月2日发的，10月底左右完成，10月底读出数据。")
        text, _ = normalize_text("公告是7月7月2日发的，十月10月底完成。", rules)
        self.assertEqual(text, "公告是7月2日发的，10月底完成。")

    def test_ambiguous_terms_are_preserved_for_review(self):
        rules = load_rules()
        text, _ = normalize_text("把基械降低百分之三十，再用踏兵。", rules)
        self.assertIn("基械", text)
        self.assertIn("踏兵", text)
        matches = ambiguous_term_matches(text, rules)
        self.assertEqual({item["observed"] for item in matches}, {"基械", "踏兵"})

    def test_apply_markdown_proposal(self):
        transcript, _, _, _ = _fixture()
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            transcript_path = root / "transcript.json"
            proposal_path = root / "review-proposal.md"
            write_json(transcript_path, transcript)
            proposal_path.write_text(
                "### RTEST S000001 00:00\n\n"
                "- 原句：专家表示行业需求增长20%，明年产能约10万吨。\n"
                "- 建议：专家表示行业需求增长20%，明年产能约12万吨。\n",
                encoding="utf-8",
            )
            result = apply_proposal_markdown(
                str(transcript_path),
                str(proposal_path),
                str(root),
            )
            self.assertEqual(result["changed_segments"], 1)
            self.assertIn("12万吨", read_json(transcript_path)["segments"][0]["text"])

    def test_channel_is_part_of_transcription_cache_config(self):
        mix = _transcription_config("mix", [])
        left = _transcription_config("0", [])
        self.assertNotEqual(sha256_json(mix), sha256_json(left))

    def test_vad_sentences_drop_punctuation_and_clean_chinese_spacing(self):
        raw = [
            {
                "sentence_info": [
                    {"text": "这 边 您 好", "start": 100, "end": 1000, "spk": 0},
                    {"text": "，", "start": 1000, "end": 1500, "spk": 0},
                    {"text": "OK 的", "start": 1500, "end": 2000, "spk": 1},
                ]
            }
        ]
        parsed = _parse_sentences(raw, 0.0, 0)
        self.assertEqual([item["text"] for item in parsed], ["这边您好", "OK 的"])
        self.assertEqual(parsed[0]["speaker"], "C01-SPK0")

    def test_short_vad_gaps_are_audited_but_not_reported(self):
        coverage = _coverage_gaps(
            [(0.0, 100.0)],
            [{"start": 0.0, "end": 40.0}, {"start": 40.8, "end": 100.0}],
        )
        self.assertEqual(coverage["status"], "pass")
        self.assertEqual(coverage["gaps"], [])
        self.assertEqual(coverage["ignored_gap_count"], 1)
        self.assertAlmostEqual(coverage["ignored_gap_seconds"], 0.8, places=3)

    def test_cached_transcription_writes_run_manifest(self):
        transcript, _, _, _ = _fixture()
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp)
            cache_root = output / "cache"
            _write_run_manifest(output, transcript, "config", cache_root, 2)
            manifest = read_json(output / "run-manifest.json")
            self.assertEqual(manifest["selected_channel"], "0")
            self.assertEqual(manifest["chunk_count"], 2)

    def test_summary_target_is_in_validator_range(self):
        transcript, minutes, _, _ = _fixture()
        count = summary_character_count(minutes["summary"])
        self.assertGreaterEqual(count, 1000)
        self.assertLessEqual(count, 1200)

    def test_validation_and_docx(self):
        transcript, minutes, review, glossary = _fixture()
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            transcript_path = root / "transcript.json"
            minutes_path = root / "minutes.json"
            review_path = root / "review.json"
            glossary_path = root / "glossary.json"
            validation_path = root / "validation.json"
            output_path = root / "会议纪要.docx"
            write_json(transcript_path, transcript)
            write_json(minutes_path, minutes)
            write_json(review_path, review)
            write_json(glossary_path, glossary)

            result = validate_minutes(
                str(minutes_path),
                str(transcript_path),
                str(review_path),
                str(validation_path),
                str(glossary_path),
            )
            self.assertTrue(result["ok"], result)
            rendered = render_docx(
                str(minutes_path),
                str(validation_path),
                str(transcript_path),
                str(review_path),
                str(output_path),
            )
            self.assertTrue(rendered.exists())

            from docx import Document

            document = Document(str(rendered))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("一、会议内容汇总", text)
            self.assertIn("二、Q&A", text)
            self.assertIn("Q1：", text)
            self.assertNotIn("\nA：", text)
            self.assertNotIn("张三", text)
            self.assertTrue(document.paragraphs[2].paragraph_format.keep_together)

    def test_open_review_blocks_validation(self):
        transcript, minutes, review, glossary = _fixture()
        review = copy.deepcopy(review)
        review["issues"][0]["status"] = "open"
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            paths = {
                "transcript": root / "transcript.json",
                "minutes": root / "minutes.json",
                "review": root / "review.json",
                "glossary": root / "glossary.json",
                "validation": root / "validation.json",
            }
            for key in ("transcript", "minutes", "review", "glossary"):
                write_json(paths[key], locals()[key])
            result = validate_minutes(
                str(paths["minutes"]),
                str(paths["transcript"]),
                str(paths["review"]),
                str(paths["validation"]),
                str(paths["glossary"]),
            )
            self.assertFalse(result["ok"])
            self.assertIn("review_open", {item["code"] for item in result["errors"]})

    def test_explicit_batch_confirmation_keeps_original_text(self):
        transcript, minutes, _, glossary = _fixture()
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            transcript_path = root / "transcript.json"
            glossary_path = root / "glossary.json"
            write_json(transcript_path, transcript)
            write_json(glossary_path, glossary)
            prepare_review(
                str(transcript_path),
                str(root),
                str(glossary_path),
            )
            prepare_review(
                str(transcript_path),
                str(root),
                str(glossary_path),
                confirm_open_as_is=True,
            )
            review = read_json(root / "review.json")
            self.assertEqual(review["status"], "ready")
            self.assertTrue(
                all(issue["status"] == "confirmed" for issue in review["issues"])
            )

    def test_prepare_review_writes_editable_proposal(self):
        transcript, _, _, glossary = _fixture()
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            transcript_path = root / "transcript.json"
            glossary_path = root / "glossary.json"
            write_json(transcript_path, transcript)
            write_json(glossary_path, glossary)
            prepare_review(str(transcript_path), str(root), str(glossary_path))
            proposal = (root / "review-proposal.md").read_text(encoding="utf-8")
            self.assertIn("- 原句：", proposal)
            self.assertIn("- 建议：", proposal)

    def test_standard_acronyms_do_not_create_review_noise(self):
        segment = {
            "id": "S000001",
            "start": 0.0,
            "end": 1.0,
            "text": "AI、DNA、mRNA 和 PD-1 是会议中的标准术语。",
        }
        issue = _review_segment_issue(
            "irrelevant",
            segment,
            {"terms": [], "person_names": []},
            load_rules(),
        )
        if issue is not None:
            self.assertNotIn("term", issue["kinds"])

    def test_review_ids_are_stable_across_transcript_revisions(self):
        segment = {
            "id": "S000001",
            "start": 0.0,
            "end": 1.0,
            "text": "MRA 肿瘤疫苗需要复核。",
        }
        first = _review_segment_issue(
            "a" * 64,
            segment,
            {"terms": [], "person_names": []},
            load_rules(),
        )
        second = _review_segment_issue(
            "b" * 64,
            segment,
            {"terms": [], "person_names": []},
            load_rules(),
        )
        self.assertIsNotNone(first)
        self.assertIsNotNone(second)
        self.assertEqual(first["id"], second["id"])

    def test_stale_review_hash_blocks_validation(self):
        transcript, minutes, review, glossary = _fixture()
        review = copy.deepcopy(review)
        review["transcript_sha256"] = "b" * 64
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            paths = {
                "transcript": root / "transcript.json",
                "minutes": root / "minutes.json",
                "review": root / "review.json",
                "glossary": root / "glossary.json",
                "validation": root / "validation.json",
            }
            write_json(paths["transcript"], transcript)
            write_json(paths["minutes"], minutes)
            write_json(paths["review"], review)
            write_json(paths["glossary"], glossary)
            result = validate_minutes(
                str(paths["minutes"]),
                str(paths["transcript"]),
                str(paths["review"]),
                str(paths["validation"]),
                str(paths["glossary"]),
            )
            self.assertFalse(result["ok"])
            self.assertIn("review_stale", {item["code"] for item in result["errors"]})

    def test_unapplied_review_replacement_blocks_validation(self):
        transcript, minutes, review, glossary = _fixture()
        review = copy.deepcopy(review)
        review["issues"][0]["replacement_text"] = "尚未应用的修正句子。"
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            paths = {
                "transcript": root / "transcript.json",
                "minutes": root / "minutes.json",
                "review": root / "review.json",
                "glossary": root / "glossary.json",
                "validation": root / "validation.json",
            }
            write_json(paths["transcript"], transcript)
            write_json(paths["minutes"], minutes)
            write_json(paths["review"], review)
            write_json(paths["glossary"], glossary)
            result = validate_minutes(
                str(paths["minutes"]),
                str(paths["transcript"]),
                str(paths["review"]),
                str(paths["validation"]),
                str(paths["glossary"]),
            )
            self.assertFalse(result["ok"])
            self.assertIn("review_not_applied", {item["code"] for item in result["errors"]})

    def test_numeric_gap_blocks_validation(self):
        transcript, minutes, review, glossary = _fixture()
        minutes = copy.deepcopy(minutes)
        minutes["summary"] = [
            paragraph.replace("20%", "两成").replace("10万吨", "十万吨")
            for paragraph in minutes["summary"]
        ]
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            paths = {
                "transcript": root / "transcript.json",
                "minutes": root / "minutes.json",
                "review": root / "review.json",
                "glossary": root / "glossary.json",
                "validation": root / "validation.json",
            }
            write_json(paths["transcript"], transcript)
            write_json(paths["minutes"], minutes)
            write_json(paths["review"], review)
            write_json(paths["glossary"], glossary)
            result = validate_minutes(
                str(paths["minutes"]),
                str(paths["transcript"]),
                str(paths["review"]),
                str(paths["validation"]),
                str(paths["glossary"]),
            )
            self.assertFalse(result["ok"])
            self.assertIn("arabic_numeral_style", {item["code"] for item in result["errors"]})

    def test_chinese_and_arabic_numerals_compare_equal(self):
        self.assertEqual(normalize_numeric_token("一点五亿"), normalize_numeric_token("1.5亿"))
        self.assertEqual(normalize_numeric_token("百分之六十"), normalize_numeric_token("60%"))
        self.assertEqual(normalize_numeric_token("三年"), normalize_numeric_token("3年"))
        self.assertEqual(normalize_numeric_token("五年半"), normalize_numeric_token("5.5年"))
        self.assertEqual(normalize_numeric_token("两个"), normalize_numeric_token("2个"))
        self.assertEqual(normalize_numeric_token("一百三四十例"), normalize_numeric_token("130-140例"))
        self.assertIn("50-60%", extract_numeric_tokens("约50-60%的收入受到影响"))

    def test_qa_prefix_is_rejected(self):
        transcript, minutes, review, glossary = _fixture()
        minutes = copy.deepcopy(minutes)
        minutes["qa"][0]["answer"][0] = "专家表示，价格可能保持稳定。"
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            paths = {
                "transcript": root / "transcript.json",
                "minutes": root / "minutes.json",
                "review": root / "review.json",
                "glossary": root / "glossary.json",
                "validation": root / "validation.json",
            }
            write_json(paths["transcript"], transcript)
            write_json(paths["minutes"], minutes)
            write_json(paths["review"], review)
            write_json(paths["glossary"], glossary)
            result = validate_minutes(
                str(paths["minutes"]),
                str(paths["transcript"]),
                str(paths["review"]),
                str(paths["validation"]),
                str(paths["glossary"]),
            )
            self.assertFalse(result["ok"])
            self.assertIn("qa_attribution_prefix", {item["code"] for item in result["errors"]})

    def test_review_apply_updates_transcript(self):
        transcript, _, review, glossary = _fixture()
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            transcript_path = root / "transcript.json"
            review_path = root / "review.json"
            glossary_path = root / "glossary.json"
            write_json(transcript_path, transcript)
            write_json(glossary_path, glossary)
            prepare_review(
                str(transcript_path),
                str(root),
                str(glossary_path),
            )
            generated_review = read_json(review_path)
            generated_review["issues"][0]["status"] = "confirmed"
            generated_review["issues"][0]["replacement_text"] = (
                "专家表示行业需求增长20%，明年产能约12万吨。"
            )
            generated_review["issues"][0]["note"] = "已回听测试"
            write_json(review_path, generated_review)
            prepare_review(
                str(transcript_path),
                str(root),
                str(glossary_path),
                apply=True,
            )
            updated = read_json(transcript_path)
            self.assertIn("12万吨", updated["segments"][0]["text"])

    def test_stale_validation_blocks_render(self):
        transcript, minutes, review, glossary = _fixture()
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            transcript_path = root / "transcript.json"
            minutes_path = root / "minutes.json"
            review_path = root / "review.json"
            glossary_path = root / "glossary.json"
            validation_path = root / "validation.json"
            output_path = root / "会议纪要.docx"
            for path, value in (
                (transcript_path, transcript),
                (minutes_path, minutes),
                (review_path, review),
                (glossary_path, glossary),
            ):
                write_json(path, value)
            result = validate_minutes(
                str(minutes_path),
                str(transcript_path),
                str(review_path),
                str(validation_path),
                str(glossary_path),
            )
            self.assertTrue(result["ok"])
            changed = copy.deepcopy(minutes)
            changed["qa"][0]["answer"][0] += "补充。"
            write_json(minutes_path, changed)
            with self.assertRaises(RenderError):
                render_docx(
                    str(minutes_path),
                    str(validation_path),
                    str(transcript_path),
                    str(review_path),
                    str(output_path),
                )


if __name__ == "__main__":
    unittest.main()
