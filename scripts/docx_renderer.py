from __future__ import annotations

from pathlib import Path
from typing import Any, Dict

from common import DEFAULT_TITLE, read_json, sha256_json


class RenderError(RuntimeError):
    pass


def _set_east_asia_font(run, font_name: str, size_pt: float, bold: bool = False) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def _configure_document(document) -> None:
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(0)

    for style_name in ("Title", "Heading 1"):
        style = document.styles[style_name]
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.color.rgb = None
        style.font.color.theme_color = None

    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def render_docx(
    minutes_path: str,
    validation_path: str,
    transcript_path: str,
    review_path: str,
    out_path: str,
    force: bool = False,
) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    minutes_file = Path(minutes_path).expanduser().resolve()
    validation_file = Path(validation_path).expanduser().resolve()
    transcript_file = Path(transcript_path).expanduser().resolve()
    review_file = Path(review_path).expanduser().resolve()
    output_file = Path(out_path).expanduser().resolve()
    if output_file.exists() and not force:
        raise RenderError("输出文件已存在；使用 --force 才会覆盖：{}".format(output_file))

    minutes = read_json(minutes_file)
    validation = read_json(validation_file)
    transcript = read_json(transcript_file)
    review = read_json(review_file)
    if not validation.get("ok"):
        raise RenderError("校验未通过，不能生成正式 DOCX。请查看 validation.json。")
    expected = {
        "minutes": sha256_json(minutes),
        "transcript": sha256_json(transcript),
        "review": sha256_json(review),
    }
    actual = {
        key: validation.get("inputs", {}).get(key, {}).get("sha256")
        for key in expected
    }
    if expected != actual:
        raise RenderError("输入文件在校验后被修改，请重新运行 validate。")
    if any(issue.get("status") != "confirmed" for issue in review.get("issues", [])):
        raise RenderError("仍有未确认待核项，不能生成正式 DOCX。")

    document = Document()
    _configure_document(document)
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = ""
    document.core_properties.subject = ""
    document.core_properties.title = DEFAULT_TITLE

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(18)
    title.paragraph_format.keep_with_next = True
    title_run = title.add_run(DEFAULT_TITLE)
    _set_east_asia_font(title_run, "黑体", 18, bold=True)

    heading = document.add_heading(level=1)
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(8)
    heading.paragraph_format.keep_with_next = True
    heading_run = heading.add_run("一、会议内容汇总")
    _set_east_asia_font(heading_run, "黑体", 14, bold=True)

    for paragraph_text in minutes["summary"]:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.first_line_indent = Pt(22)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.widow_control = True
        run = paragraph.add_run(paragraph_text.strip())
        _set_east_asia_font(run, "宋体", 11)

    heading = document.add_heading(level=1)
    heading.paragraph_format.space_before = Pt(16)
    heading.paragraph_format.space_after = Pt(8)
    heading.paragraph_format.keep_with_next = True
    heading_run = heading.add_run("二、Q&A")
    _set_east_asia_font(heading_run, "黑体", 14, bold=True)

    for item in minutes["qa"]:
        question = document.add_paragraph()
        question.paragraph_format.space_before = Pt(6)
        question.paragraph_format.space_after = Pt(4)
        question.paragraph_format.keep_with_next = True
        label = question.add_run("{}：".format(item["id"]))
        _set_east_asia_font(label, "宋体", 11, bold=True)
        question_text = question.add_run(str(item["question"]).strip())
        _set_east_asia_font(question_text, "宋体", 11)
        for answer_text in item["answer"]:
            answer = document.add_paragraph()
            answer.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            answer.paragraph_format.first_line_indent = Pt(22)
            answer.paragraph_format.space_after = Pt(4)
            answer.paragraph_format.keep_together = True
            answer.paragraph_format.widow_control = True
            run = answer.add_run(str(answer_text).strip())
            _set_east_asia_font(run, "宋体", 11)

    output_file.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_file.with_suffix(".tmp.docx")
    document.save(str(temporary))
    temporary.replace(output_file)
    return output_file
